from __future__ import annotations

from shutil import copyfile
from pathlib import Path
from io import StringIO
import csv
import json

from openpyxl import Workbook
import pytest


def _books(tmp_path):
    first = Workbook()
    positive = first.active
    positive.title = "重复"
    positive.append(["改写ID", "原文", "重复组", "原标注备注", "最终判定"])
    positive.append(["101", "甲公司宣布回购。", "1", None, "重复"])
    positive.append(["102", "甲公司完成回购。", "1", None, "重复"])
    positive.append(["103", "乙公司宣布回购。", "2", None, "不重复"])
    negative = first.create_sheet("不重复")
    negative.append(["改写ID_A", "改写ID_B", "文本A", "文本B", "类型", "最终判定"])
    negative.append(["102", "103", "甲公司完成回购。", "乙公司宣布回购。", "主体", "不重复"])
    negative.row_dimensions[2].hidden = True
    first_path = tmp_path / "first.xlsx"
    first.save(first_path)

    second = Workbook()
    labeled = second.active
    labeled.title = "第二批_会议口径标注"
    labeled.append(["改写ID", "原文", "来源", "重复组", "会议口径最终判定"])
    labeled.append(["201", "丙公司宣布回购。", "来源甲", "1", "重复"])
    labeled.append(["202", "丙公司完成回购。", "来源乙", "1", "重复"])
    second_path = tmp_path / "second.xlsx"
    second.save(second_path)
    return first_path, second_path


def test_load_workbooks_keeps_hidden_endpoints_rows_and_legacy_labels(tmp_path):
    from news_flash_dedup.gold import load_workbooks

    catalog = load_workbooks(_books(tmp_path))

    assert len(catalog.materials) == 2
    assert len(catalog.rows) == 7
    assert len(catalog.negative_pairs) == 1
    hidden = [row for row in catalog.rows if row.sheet_name == "不重复"]
    assert len(hidden) == 2
    assert {row.side for row in hidden} == {"A", "B"}
    assert all(row.hidden and row.excel_row == 2 for row in hidden)
    assert {row.legacy_label for row in hidden} == {"不重复"}
    assert sum(row.legacy_label == "不重复" for row in catalog.rows) == 3
    assert all(not hasattr(row, "gold_label") for row in catalog.rows)


def test_t072_group_and_negative_pair_edges_keep_every_endpoint_in_one_split(tmp_path):
    from news_flash_dedup.gold import assign_splits, build_families, load_workbooks

    catalog = load_workbooks(_books(tmp_path))
    graph = build_families(catalog)
    split = assign_splits(graph, seed="frozen-v1")
    first = [row for row in catalog.rows if row.sheet_name in {"重复", "不重复"}]
    second = [row for row in catalog.rows if row.sheet_name == "第二批_会议口径标注"]

    assert len({graph.family_by_row[row.row_key] for row in first}) == 1
    assert len({graph.family_by_row[row.row_key] for row in second}) == 1
    assert graph.family_by_row[first[0].row_key] != graph.family_by_row[second[0].row_key]
    assert len({split[graph.family_by_row[row.row_key]] for row in first}) == 1
    assert len({split[graph.family_by_row[row.row_key]] for row in second}) == 1
    assert set(split.values()) <= {"dev", "validation", "test"}
    assert "negative_pair" in {edge.reason for edge in graph.edges}


def test_t072_identical_file_alias_and_cross_sheet_identity_risks_are_visible(tmp_path):
    from news_flash_dedup.gold import build_families, load_workbooks

    first, second = _books(tmp_path)
    from openpyxl import load_workbook

    workbook = load_workbook(first)
    workbook["不重复"].append([
        "101", "900", "甲公司修订回购。", "甲公司宣布回购。", "正文", "不重复",
    ])
    workbook.save(first)
    alias = tmp_path / "first-copy.xlsx"
    copyfile(first, alias)

    catalog = load_workbooks((alias, second, first))
    graph = build_families(catalog)
    assert len(catalog.materials) == 2
    assert len(catalog.rows) == 9
    assert len(catalog.negative_pairs) == 2
    assert {risk.kind for risk in graph.risks} >= {"id_body_conflict", "body_id_alias"}
    assert all(risk.negative_row == 3 for risk in graph.risks)
    assert len({graph.family_by_row[row.row_key] for row in catalog.rows if row.material_id == catalog.negative_pairs[1].material_id}) == 1


def test_t072_documented_cross_group_leads_only_connect_isolation_family(tmp_path):
    from news_flash_dedup.gold import build_families, load_workbooks

    first, second = _books(tmp_path)
    from openpyxl import load_workbook

    workbook = load_workbook(first)
    positive = workbook["重复"]
    positive.append(["301", "丁公司发行债券。", "25", None, "重复"])
    positive.append(["302", "戊公司收购资产。", "27", None, "重复"])
    workbook.save(first)

    catalog = load_workbooks((first, second))
    graph = build_families(catalog)
    leads = [row for row in catalog.rows if row.legacy_group in {"25", "27"}]
    assert len(leads) == 2
    assert graph.family_by_row[leads[0].row_key] == graph.family_by_row[leads[1].row_key]
    assert any(edge.reason == "documented_cross_group_lead" for edge in graph.edges)
    assert all(not hasattr(row, "gold_label") for row in leads)


def test_manifest_and_blank_double_annotation_template_never_export_body_or_gold(tmp_path):
    from news_flash_dedup.gold import (
        assign_splits, blank_annotation_csv, build_families, build_manifest, load_workbooks,
    )

    first, second = _books(tmp_path)
    catalog = load_workbooks((first, second))
    graph = build_families(catalog)
    manifest = build_manifest(
        catalog, graph, assign_splits(graph, seed="frozen-v1"), split_seed="frozen-v1",
    )
    exported = json.dumps(manifest, ensure_ascii=False)
    left, right = catalog.rows[:2]
    csv_text = blank_annotation_csv(
        catalog, ((left.row_key, right.row_key),), rule_version="P03-draft-2026-09-23",
    )

    assert manifest["totals"]["row_records"] == 7
    assert manifest["totals"]["negative_pair_rows"] == 1
    assert manifest["totals"]["hidden_row_nodes"] == 2
    assert manifest["split_seed"] == "frozen-v1"
    assert len(manifest["rows"]) == 7
    assert all(len(item["raw_hash"]) == 64 for item in manifest["rows"])
    assert sum(item["hidden"] for item in manifest["rows"]) == 2
    assert manifest["edge_counts"]["negative_pair"] == 1
    assert "raw_text" not in exported
    assert "甲公司宣布回购" not in exported + csv_text
    assert "gold_label" not in exported + csv_text
    assert "legacy_label" not in csv_text
    tasks = list(csv.DictReader(StringIO(csv_text)))
    assert len(tasks) == 2
    assert {task["annotator_slot"] for task in tasks} == {"1", "2"}
    assert all(task["decision"] == "" and task["annotation_status"] == "pending" for task in tasks)
    assert all(task["rule_version"] == "P03-draft-2026-09-23" for task in tasks)
    assert all(task["annotator_id"] == "" and task["coverage_complete"] == "" for task in tasks)
    reversed_csv = blank_annotation_csv(
        catalog, ((right.row_key, left.row_key),), rule_version="P03-draft-2026-09-23",
    )
    assert next(csv.DictReader(StringIO(reversed_csv)))["pair_id"] == tasks[0]["pair_id"]
    with pytest.raises(ValueError, match="more than once"):
        blank_annotation_csv(
            catalog,
            ((left.row_key, right.row_key), (right.row_key, left.row_key)),
            rule_version="P03-draft-2026-09-23",
        )
    assert manifest["families"]


def test_manifest_and_family_split_do_not_depend_on_input_path_order(tmp_path):
    from news_flash_dedup.gold import assign_splits, build_families, build_manifest, load_workbooks

    first, second = _books(tmp_path)
    direct = load_workbooks((first, second))
    reversed_input = load_workbooks((second, first))
    graph_a, graph_b = build_families(direct), build_families(reversed_input)
    manifest_a = build_manifest(
        direct, graph_a, assign_splits(graph_a, seed="frozen-v1"), split_seed="frozen-v1",
    )
    manifest_b = build_manifest(
        reversed_input, graph_b, assign_splits(graph_b, seed="frozen-v1"),
        split_seed="frozen-v1",
    )
    assert manifest_a == manifest_b


def test_t072_real_materials_preserve_hidden_rows_and_never_split_a_family():
    from news_flash_dedup.gold import assign_splits, build_families, load_workbooks

    root = Path(__file__).resolve().parents[4]
    first = root / "v2_重复case_Fact特征分类标注版_与用户标注一致版.xlsx"
    second = root / "v2_重复case_第二批_会议口径标注版.xlsx"
    if not first.exists() or not second.exists():
        pytest.skip("历史材料未随代码提供")

    catalog = load_workbooks((first, second))
    graph = build_families(catalog)
    split = assign_splits(graph, seed="p04-material-v1")
    assert len(catalog.rows) == 1035
    assert len(catalog.negative_pairs) == 42
    assert sum(row.hidden for row in catalog.rows) == 18
    assert all(
        graph.family_by_row[pair.a_row_key] == graph.family_by_row[pair.b_row_key]
        for pair in catalog.negative_pairs
    )
    assert all(
        len({split[graph.family_by_row[key]] for key in members}) == 1
        for members in graph.members.values()
    )
    first_batch = [row for row in catalog.rows if row.sheet_name == "重复"]
    for a_group, b_group in (("25", "27"), ("301", "303"), ("375", "376"), ("292", "294"), ("306", "307")):
        a_rows = [row for row in first_batch if row.legacy_group == a_group]
        b_rows = [row for row in first_batch if row.legacy_group == b_group]
        assert a_rows and b_rows
        assert {graph.family_by_row[row.row_key] for row in a_rows + b_rows} == {
            graph.family_by_row[a_rows[0].row_key]
        }
    rows_by_location = {(row.sheet_name, row.excel_row, row.side): row for row in catalog.rows}
    for left, right in ((350, 351), (354, 355)):
        a = rows_by_location[("重复", left, "single")]
        b = rows_by_location[("重复", right, "single")]
        assert a.raw_id == b.raw_id
        assert graph.family_by_row[a.row_key] == graph.family_by_row[b.row_key]
    a = rows_by_location[("第二批_会议口径标注", 44, "single")]
    b = rows_by_location[("第二批_会议口径标注", 45, "single")]
    assert a.raw_text == b.raw_text and a.raw_id != b.raw_id
    assert graph.family_by_row[a.row_key] == graph.family_by_row[b.row_key]
