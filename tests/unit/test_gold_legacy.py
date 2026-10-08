"""历史修订表的显式成对标签；不把多行组扩为 pair。"""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from news_flash_dedup.gold import assign_splits, build_families, load_workbooks


def _books(tmp_path):
    first = Workbook()
    sheet = first.active
    sheet.title = "重复"
    sheet.append(["改写ID", "原文", "重复组", "原标注备注", "最终判定"])
    for row in (
        ("1", "fixture-legacy-alpha", "1", None, "重复"),
        ("2", "fixture-legacy-beta", "1", None, "重复"),
        ("3", "fixture-legacy-gamma", "2", None, "重复"),
        ("4", "fixture-legacy-delta", "2", None, "重复"),
        ("5", "fixture-legacy-epsilon", "2", None, "重复"),
        ("6", "fixture-legacy-zeta", "3", None, "不重复"),
        ("7", "fixture-legacy-eta", "3", None, "不重复"),
        ("8", "fixture-legacy-theta", "50", None, "重复"),
        ("9", "fixture-legacy-iota", "50", None, "重复"),
    ):
        sheet.append(row)
    negative = first.create_sheet("不重复")
    negative.append(["改写ID_A", "改写ID_B", "文本A", "文本B", "类型", "最终判定"])
    for _ in (1, 2):
        negative.append(["10", "11", "fixture-legacy-kappa", "fixture-legacy-lambda", "主体", "不重复"])
    first_path = tmp_path / "first.xlsx"
    first.save(first_path)

    second = Workbook()
    sheet = second.active
    sheet.title = "第二批_会议口径标注"
    sheet.append(["改写ID", "原文", "来源", "重复组", "会议口径最终判定"])
    for i, group in enumerate(("1", "1", "50", "50", "50", "50"), start=20):
        sheet.append([str(i), f"fixture-legacy-second-{i}", "source fixture", group, "重复"])
    second_path = tmp_path / "second.xlsx"
    second.save(second_path)
    return first_path, second_path


def _freeze(first, second):
    from news_flash_dedup.gold.legacy import freeze_legacy_reviewed_initial

    catalog = load_workbooks((first, second))
    graph = build_families(catalog)
    splits = assign_splits(graph, seed="fixture-legacy-split")
    return freeze_legacy_reviewed_initial(
        catalog=catalog,
        graph=graph,
        splits=splits,
        split_seed="fixture-legacy-split",
        first_workbook=first,
        second_workbook=second,
        frozen_at=datetime(2026, 9, 24, tzinfo=timezone.utc),
        synthetic_plan_id="fixture-source-row-order-v1",
        synthetic_business_date="2000-01-01",
    )


def test_legacy_initial_only_freezes_explicit_two_row_and_negative_pairs(tmp_path):
    first, second = _books(tmp_path)
    frozen = _freeze(first, second)
    manifest = frozen.manifest

    assert manifest["provenance_type"] == "legacy_reviewed"
    assert manifest["counts"]["source_pair_records"] == 6
    assert manifest["counts"]["unique_explicit_pairs"] == 5
    assert manifest["counts"]["clear_pairs"] == 4
    assert manifest["counts"]["held_pairs"] == 1
    assert manifest["counts"]["uncovered_multiline_groups"] == 2
    assert manifest["counts"]["synthetic_replay_rows"] == 15
    assert manifest["counts"]["M"] == 0 and manifest["counts"]["H"] == 0
    assert all(item["evidence_gap"] == "no_original_span_evidence" for item in manifest["pairs"])
    assert all(item["annotator_ids"] is None for item in manifest["pairs"])
    assert all(item["adjudicator_ids"] is None for item in manifest["pairs"])
    assert all(row.raw_text not in frozen.manifest_json for row in load_workbooks((first, second)).rows)
    assert manifest["synthetic_plan"]["order_provenance"] == "synthetic"
    assert manifest["synthetic_plan"]["business_date"] == "2000-01-01"
    assert frozen.dataset_version.startswith("ds-legacy-v1-")


def test_legacy_hold_preserves_source_label_and_no_group_transmission(tmp_path):
    first, second = _books(tmp_path)
    frozen = _freeze(first, second)
    pairs = frozen.manifest["pairs"]
    held = [item for item in pairs if item["status"] == "hold"]
    assert len(held) == 1
    assert held[0]["source_final_label"] == "重复"
    assert held[0]["gold_label"] is None
    assert held[0]["hold_reason"] == "numeric_zero_tolerance_review"
    assert {source["excel_row"] for source in held[0]["sources"]} == {9, 10}
    assert all(item["source_group"] != "2" for item in pairs)
    assert any(item["source_group"] == "2" for item in frozen.manifest["uncovered_groups"])
    assert any(item["source_group"] == "50" for item in frozen.manifest["uncovered_groups"])


def test_legacy_mixed_final_columns_are_held_and_change_dataset_version(tmp_path):
    first, second = _books(tmp_path)
    original = _freeze(first, second)
    workbook = load_workbook(first)
    workbook["重复"].cell(3, 5).value = "不重复"
    workbook.save(first)
    changed = _freeze(first, second)
    assert changed.dataset_version != original.dataset_version
    assert changed.manifest["counts"]["held_pairs"] == 2
    mixed = [item for item in changed.manifest["pairs"] if item["hold_reason"] == "mixed_final_labels"]
    assert len(mixed) == 1
    assert set(mixed[0]["source_final_values"]) == {"重复", "不重复"}


def test_legacy_missing_final_column_is_held_without_invented_label(tmp_path):
    first, second = _books(tmp_path)
    workbook = load_workbook(first)
    workbook["重复"].cell(3, 5).value = None
    workbook.save(first)
    frozen = _freeze(first, second)
    held = [item for item in frozen.manifest["pairs"] if item["hold_reason"] == "mixed_final_labels"]
    assert len(held) == 1
    assert held[0]["gold_label"] is None
    assert held[0]["source_final_values"] == [None, "重复"]


def test_legacy_real_materials_pair_counts_and_hold_rows():
    root = Path(__file__).resolve().parents[4]
    first = root / "v2_重复case_Fact特征分类标注版_与用户标注一致版.xlsx"
    second = root / "v2_重复case_第二批_会议口径标注版.xlsx"
    if not first.exists() or not second.exists():
        pytest.skip("historical workbooks are not distributed with package")
    frozen = _freeze(first, second)
    counts = frozen.manifest["counts"]
    assert counts["source_pair_records"] == 395
    assert counts["unique_explicit_pairs"] == 393
    assert counts["clear_pairs"] == 389
    assert counts["held_pairs"] == 4
    assert counts["uncovered_multiline_groups"] == 72
    assert counts["uncovered_group_rows"] == 245
    assert counts["synthetic_replay_rows"] == 951
    holds = {
        item["source_group"]: {source["excel_row"] for source in item["sources"]}
        for item in frozen.manifest["pairs"] if item["status"] == "hold"
    }
    assert holds == {"6": {12, 13}, "29": {63, 64}, "50": {113, 114}, "275": {568, 569}}
