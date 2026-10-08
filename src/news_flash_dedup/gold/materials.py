"""只读历史材料解析，保留原始行身份和旧标签。"""

from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


@dataclass(frozen=True)
class MaterialFile:
    material_id: str
    byte_size: int
    aliases: tuple[str, ...]
    sheet_counts: tuple[tuple[str, int], ...]


@dataclass(frozen=True)
class MaterialRow:
    row_key: str
    material_id: str
    sheet_name: str
    excel_row: int
    side: str
    raw_id: str
    raw_text: str
    raw_hash: str
    legacy_group: str | None
    legacy_label: str | None
    legacy_note: str | None
    hidden: bool


@dataclass(frozen=True)
class NegativePair:
    material_id: str
    sheet_name: str
    excel_row: int
    a_row_key: str
    b_row_key: str
    legacy_label: str | None


@dataclass(frozen=True)
class MaterialCatalog:
    materials: tuple[MaterialFile, ...]
    rows: tuple[MaterialRow, ...]
    negative_pairs: tuple[NegativePair, ...]


def _file_hash(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _value(sheet, row: int, column: int) -> str | None:
    value = sheet.cell(row, column).value
    return None if value is None else str(value)


def _make_row(
    material_id: str,
    sheet,
    row_number: int,
    side: str,
    raw_id: str | None,
    raw_text: str | None,
    legacy_group: str | None,
    legacy_label: str | None,
    legacy_note: str | None,
) -> MaterialRow:
    if not raw_id or raw_text is None or raw_text == "":
        raise ValueError(f"material row lacks ID or body: {sheet.title}:{row_number}:{side}")
    return MaterialRow(
        row_key=f"{material_id}:{sheet.title}:{row_number}:{side}",
        material_id=material_id,
        sheet_name=sheet.title,
        excel_row=row_number,
        side=side,
        raw_id=raw_id,
        raw_text=raw_text,
        raw_hash=sha256(raw_text.encode("utf-8")).hexdigest(),
        legacy_group=legacy_group,
        legacy_label=legacy_label,
        legacy_note=legacy_note,
        hidden=bool(sheet.row_dimensions[row_number].hidden),
    )


def load_workbooks(paths: Iterable[str | Path]) -> MaterialCatalog:
    """读取两批历史表；相同字节的路径只计一份材料。"""
    by_hash: dict[str, list[Path]] = {}
    for supplied in paths:
        path = Path(supplied).resolve(strict=True)
        digest = _file_hash(path)
        by_hash.setdefault(digest, []).append(path)

    materials: list[MaterialFile] = []
    rows: list[MaterialRow] = []
    negative_pairs: list[NegativePair] = []
    for material_id, aliases in sorted(by_hash.items()):
        source = min(aliases, key=lambda path: str(path))
        workbook = load_workbook(source, read_only=False, data_only=True)
        sheet_counts: list[tuple[str, int]] = []
        try:
            for sheet in workbook.worksheets:
                if sheet.title not in {"重复", "不重复", "第二批_会议口径标注"}:
                    sheet_counts.append((sheet.title, 0))
                    continue
                count = 0
                for row_number in range(2, sheet.max_row + 1):
                    if sheet.title == "不重复":
                        label = _value(sheet, row_number, 6)
                        left = _make_row(
                            material_id, sheet, row_number, "A",
                            _value(sheet, row_number, 1), _value(sheet, row_number, 3),
                            None, label, _value(sheet, row_number, 5),
                        )
                        right = _make_row(
                            material_id, sheet, row_number, "B",
                            _value(sheet, row_number, 2), _value(sheet, row_number, 4),
                            None, label, _value(sheet, row_number, 5),
                        )
                        rows.extend((left, right))
                        negative_pairs.append(NegativePair(
                            material_id, sheet.title, row_number,
                            left.row_key, right.row_key, label,
                        ))
                    else:
                        first_batch = sheet.title == "重复"
                        rows.append(_make_row(
                            material_id, sheet, row_number, "single",
                            _value(sheet, row_number, 1), _value(sheet, row_number, 2),
                            _value(sheet, row_number, 3 if first_batch else 4),
                            _value(sheet, row_number, 5),
                            _value(sheet, row_number, 4 if first_batch else 23),
                        ))
                    count += 1
                sheet_counts.append((sheet.title, count))
        finally:
            workbook.close()
        materials.append(MaterialFile(
            material_id=material_id,
            byte_size=source.stat().st_size,
            aliases=tuple(sorted(str(path) for path in aliases)),
            sheet_counts=tuple(sheet_counts),
        ))
    return MaterialCatalog(tuple(materials), tuple(rows), tuple(negative_pairs))
