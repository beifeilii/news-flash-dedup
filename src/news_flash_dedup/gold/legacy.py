"""两本用户裁定修订表的金标初版冻结；只采显式对，不扩展多行组。"""

from __future__ import annotations

import json
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

from .artifacts import build_manifest
from .families import FamilyGraph
from .materials import MaterialCatalog, MaterialRow


FINAL_LABELS = frozenset({"重复", "不重复"})
FIRST_BATCH_HOLDS = {
    "6": "multi_event_relation_review",
    "29": "multi_event_historical_reference_review",
    "50": "numeric_zero_tolerance_review",
    "275": "subject_conflict_review",
}


def _canonical(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _hash(value: str) -> str:
    return sha256(value.encode("utf-8")).hexdigest()


def _file_hash(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _value(sheet, row: int, col: int) -> str | None:
    value = sheet.cell(row, col).value
    return None if value is None else str(value)


def _reason_hash(sheet, row: int, columns: tuple[int, ...]) -> str | None:
    reasons = [value for col in columns if (value := _value(sheet, row, col))]
    return _hash(_canonical(reasons)) if reasons else None


def _identity(row: MaterialRow) -> tuple[str, str]:
    return row.raw_id, row.raw_hash


def _identity_pair(a: MaterialRow, b: MaterialRow) -> tuple[tuple[str, str], tuple[str, str]]:
    return tuple(sorted((_identity(a), _identity(b))))


@dataclass(frozen=True)
class FrozenLegacyDataset:
    dataset_version: str
    manifest_json: str

    @property
    def manifest(self) -> dict:
        return json.loads(self.manifest_json)


def freeze_legacy_reviewed_initial(
    *,
    catalog: MaterialCatalog,
    graph: FamilyGraph,
    splits: dict[str, str],
    split_seed: str,
    first_workbook: str | Path,
    second_workbook: str | Path,
    frozen_at: datetime,
    synthetic_plan_id: str,
    synthetic_business_date: str,
) -> FrozenLegacyDataset:
    """仅冻结两本已获用户裁定表中的一对一来源关系及无正文行序计划。"""
    if frozen_at.tzinfo is None or frozen_at.utcoffset() is None:
        raise ValueError("frozen_at needs a timezone")
    if not synthetic_plan_id.strip():
        raise ValueError("synthetic plan ID is required")
    try:
        parsed_date = date.fromisoformat(synthetic_business_date)
    except ValueError as exc:
        raise ValueError("synthetic business date must be ISO") from exc
    if parsed_date.isoformat() != synthetic_business_date:
        raise ValueError("synthetic business date must be canonical ISO")

    isolation = build_manifest(catalog, graph, splits, split_seed=split_seed)
    material_ids = {material.material_id for material in catalog.materials}
    first_path = Path(first_workbook).resolve(strict=True)
    second_path = Path(second_workbook).resolve(strict=True)
    first_id, second_id = _file_hash(first_path), _file_hash(second_path)
    if first_id == second_id or {first_id, second_id} != material_ids or len(material_ids) != 2:
        raise ValueError("workbook bytes must match exactly the two catalog materials")

    rows_by_key = {row.row_key: row for row in catalog.rows}
    grouped: dict[tuple[str, str], list[MaterialRow]] = defaultdict(list)
    for row in catalog.rows:
        if row.side == "single" and row.legacy_group is not None:
            grouped[(row.material_id, row.legacy_group)].append(row)

    workbooks = {
        first_id: load_workbook(first_path, read_only=False, data_only=True),
        second_id: load_workbook(second_path, read_only=False, data_only=True),
    }
    candidates: dict[tuple[tuple[str, str], tuple[str, str]], list[dict]] = defaultdict(list)
    uncovered: list[dict] = []
    synthetic_rows: list[dict] = []
    try:
        for material_id, sheet_name, group_col, reason_cols in (
            (first_id, "重复", 3, (12, 13)),
            (second_id, "第二批_会议口径标注", 4, (20, 21)),
        ):
            sheet = workbooks[material_id][sheet_name]
            sheet_rows = sorted(
                (row for row in catalog.rows if row.material_id == material_id and row.sheet_name == sheet_name),
                key=lambda row: row.excel_row,
            )
            for index, row in enumerate(sheet_rows, start=1):
                synthetic_rows.append({
                    "row_key": row.row_key,
                    "scope_id": f"synthetic:{material_id[:16]}",
                    "business_date": synthetic_business_date,
                    "arrival_seq": index,
                })
            groups = {
                group: sorted(members, key=lambda row: row.excel_row)
                for (source, group), members in grouped.items() if source == material_id
            }
            for group, members in sorted(groups.items()):
                if len(members) != 2:
                    uncovered.append({
                        "material_id": material_id,
                        "sheet_name": sheet_name,
                        "source_group": group,
                        "row_keys": [row.row_key for row in members],
                        "row_count": len(members),
                        "final_values": sorted(
                            {row.legacy_label for row in members},
                            key=lambda value: "" if value is None else value,
                        ),
                        "reason": (
                            "multirow_group_with_B133_missing_predicate"
                            if material_id == second_id and group == "50"
                            else "multirow_group_pair_undetermined"
                        ),
                    })
                    continue
                a, b = members
                if _value(sheet, a.excel_row, group_col) != group or _value(sheet, b.excel_row, group_col) != group:
                    raise ValueError("catalog group does not match workbook")
                sources = [
                    {
                        "material_id": material_id,
                        "file_name": first_path.name if material_id == first_id else second_path.name,
                        "sheet_name": sheet_name,
                        "excel_row": row.excel_row,
                        "row_key": row.row_key,
                        "final_column": "E",
                        "final_value": _value(sheet, row.excel_row, 5),
                        "reason_sha256": _reason_hash(sheet, row.excel_row, reason_cols),
                    }
                    for row in members
                ]
                candidates[_identity_pair(a, b)].append({
                    "source_kind": "two_row_group",
                    "material_id": material_id,
                    "sheet_name": sheet_name,
                    "source_group": group,
                    "sources": sources,
                    "endpoint_row_keys": [a.row_key, b.row_key],
                })

        negative = workbooks[first_id]["不重复"]
        for pair in catalog.negative_pairs:
            if pair.material_id != first_id:
                raise ValueError("negative pair must belong to first workbook")
            a, b = rows_by_key[pair.a_row_key], rows_by_key[pair.b_row_key]
            candidates[_identity_pair(a, b)].append({
                "source_kind": "explicit_negative_sheet",
                "material_id": first_id,
                "sheet_name": "不重复",
                "source_group": None,
                "sources": [{
                    "material_id": first_id,
                    "file_name": first_path.name,
                    "sheet_name": "不重复",
                    "excel_row": pair.excel_row,
                    "row_key": pair.a_row_key,
                    "other_row_key": pair.b_row_key,
                    "final_column": "F",
                    "final_value": _value(negative, pair.excel_row, 6),
                    "reason_sha256": _reason_hash(negative, pair.excel_row, (13, 14)),
                }],
                "endpoint_row_keys": [a.row_key, b.row_key],
            })
    finally:
        for workbook in workbooks.values():
            workbook.close()

    pairs: list[dict] = []
    for identities, origins in sorted(candidates.items()):
        all_sources = [source for origin in origins for source in origin["sources"]]
        values = sorted(
            {source["final_value"] for source in all_sources},
            key=lambda value: "" if value is None else value,
        )
        row_keys = sorted({key for origin in origins for key in origin["endpoint_row_keys"]})
        family_ids = {graph.family_by_row[key] for key in row_keys}
        if len(family_ids) != 1:
            raise ValueError("one explicit pair crossed isolation families")
        source_groups = {origin["source_group"] for origin in origins if origin["source_group"] is not None}
        source_group = next(iter(source_groups)) if len(source_groups) == 1 else None
        hold_reason = None
        if len(values) != 1:
            hold_reason = "mixed_final_labels"
        elif values[0] not in FINAL_LABELS:
            hold_reason = "invalid_final_label"
        elif any(origin["material_id"] == first_id and origin["sheet_name"] == "重复" for origin in origins):
            hold_reason = FIRST_BATCH_HOLDS.get(source_group)
        source_label = values[0] if len(values) == 1 else None
        family_id = next(iter(family_ids))
        pairs.append({
            "pair_id": _hash(_canonical(identities)),
            "provenance_type": "legacy_reviewed",
            "status": "hold" if hold_reason else "clear",
            "gold_label": None if hold_reason else source_label,
            "source_final_label": source_label,
            "source_final_values": values,
            "hold_reason": hold_reason,
            "source_group": source_group,
            "family_id": family_id,
            "split": splits[family_id],
            "endpoint_row_keys": row_keys,
            "sources": sorted(all_sources, key=lambda item: (item["material_id"], item["sheet_name"], item["excel_row"])),
            "evidence_gap": "no_original_span_evidence",
            "annotator_ids": None,
            "adjudicator_ids": None,
        })

    pairs.sort(key=lambda item: item["pair_id"])
    uncovered.sort(key=lambda item: (item["material_id"], item["sheet_name"], item["source_group"]))
    synthetic_rows.sort(key=lambda item: (item["scope_id"], item["arrival_seq"]))
    label_counts = Counter(item["gold_label"] for item in pairs if item["status"] == "clear")
    manifest = {
        "schema_version": "legacy-v1",
        "provenance_type": "legacy_reviewed",
        "source_authorization": "2026-09-23 user ruling in 07 §7 and 12 §4.2",
        "frozen_at": frozen_at.isoformat(),
        "isolation_manifest_sha256": _hash(_canonical(isolation)),
        "split_seed": split_seed,
        "materials": [
            {"material_id": material.material_id, "byte_size": material.byte_size, "file_names": sorted({Path(path).name for path in material.aliases})}
            for material in catalog.materials
        ],
        "pairs": pairs,
        "uncovered_groups": uncovered,
        "synthetic_plan": {
            "plan_id": synthetic_plan_id,
            "order_provenance": "synthetic",
            "business_date": synthetic_business_date,
            "scope_policy": "one artificial scope per workbook; not a real business scope",
            "order_policy": "positive-sheet Excel row order; negative-sheet pair rows excluded from stream",
            "metric_status": "synthetic_only_no_real_M_H_or_online_recall",
            "rows": synthetic_rows,
        },
        "counts": {
            "source_pair_records": sum(
                1 for origins in candidates.values() for _ in origins
            ),
            "unique_explicit_pairs": len(pairs),
            "clear_pairs": sum(item["status"] == "clear" for item in pairs),
            "clear_duplicate_pairs": label_counts["重复"],
            "clear_nonduplicate_pairs": label_counts["不重复"],
            "held_pairs": sum(item["status"] == "hold" for item in pairs),
            "uncovered_multiline_groups": len(uncovered),
            "uncovered_group_rows": sum(item["row_count"] for item in uncovered),
            "synthetic_replay_rows": len(synthetic_rows),
            "M": 0,
            "H": 0,
            "q_plus": 0,
        },
    }
    version = "ds-legacy-v1-" + _hash(_canonical(manifest))
    manifest["dataset_version"] = version
    return FrozenLegacyDataset(version, _canonical(manifest))
